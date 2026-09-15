import os
from collections import Counter
from datetime import date, datetime

import frappe
from frappe import _
from frappe.model.document import Document
from frappe.utils import flt, now, nowdate

from employee_lending.employee_lending.doctype.employee_lending_settings.employee_lending_settings import get_settings
from employee_lending.employee_lending.legacy_import import (
    aggregate_gl_rows,
    classify_outstanding_row,
    map_gl_headers,
    map_headers,
)


class EmployeeLoanLegacyImportBatch(Document):
    def before_insert(self):
        self.status = "Draft"
        if not self.cutoff_date:
            self.cutoff_date = nowdate()

    def validate(self):
        settings = get_settings()
        self.company = settings.company
        if self.docstatus == 0 and self.status not in (None, "", "Draft"):
            frappe.throw(_("A processed batch cannot be edited"))

    @frappe.whitelist()
    def load_source_file(self):
        self.check_permission("write")
        if self.docstatus != 0:
            frappe.throw(_("Only a draft batch can load a report"))
        if not self.source_file:
            frappe.throw(_("Attach the Employee Loan Outstanding XLSX report"))
        if not self.gl_source_file:
            frappe.throw(_("Attach the production General Ledger XLSX report"))
        if not str(self.source_file).lower().endswith(".xlsx") or not str(self.gl_source_file).lower().endswith(".xlsx"):
            frappe.throw(_("Both source files must be XLSX workbooks"))

        settings = get_settings()
        records = read_outstanding_file(self.source_file)
        gl_by_employee = read_gl_file(self.gl_source_file, settings, self.cutoff_date)
        employee_ids = sorted({record["employee"] for _, record in records if record})
        employees = {
            row.name: row
            for row in frappe.get_all(
                "Employee",
                filters={"name": ["in", employee_ids]},
                fields=["name", "employee_name", "company"],
                limit_page_length=0,
            )
        }
        products = frappe.get_all(
            "Employee Loan Product",
            filters={"disabled": 0},
            fields=["name", "flat_interest_rate", "number_of_fortnights"],
            order_by="flat_interest_rate, name",
            limit_page_length=0,
        )

        employee_counts = Counter(record["employee"] for _, record in records if record)
        self.set("rows", [])
        for row_number, record in records:
            if not record:
                continue
            status = record["validation_status"]
            message = record["validation_message"]
            employee = employees.get(record["employee"])
            product = match_product(record["source_interest_rate"], products)
            employee_gl = gl_by_employee.get(record["employee"])

            if employee_counts[record["employee"]] > 1:
                status, message = "Invalid", "Duplicate employee in source report"
            elif not employee:
                status, message = "Invalid", "Employee does not exist in ERPNext"
            elif employee.company != self.company:
                status, message = "Invalid", "Employee company does not match the batch company"
            elif status == "Ready" and not product:
                status, message = "Invalid", "No unique active loan product matches the source interest rate"
            elif status == "Ready":
                status, message = validate_gl_reconciliation(record, employee_gl)

            self.append(
                "rows",
                {
                    "row_number": row_number,
                    "source_employee_id": record["employee"],
                    "employee": employee.name if employee else None,
                    "employee_name": employee.employee_name if employee else "",
                    "loan_product": product.name if product else "",
                    "source_interest_rate": float(record["source_interest_rate"]),
                    "original_principal": float(record["original_principal"]),
                    "original_interest": float(record["original_interest"]),
                    "principal_repaid": float(record["principal_repaid"]),
                    "interest_repaid": float(record["interest_repaid"]),
                    "principal_outstanding": float(record["principal_outstanding"]),
                    "interest_outstanding": float(record["interest_outstanding"]),
                    "reported_total_outstanding": float(record["reported_total_outstanding"]),
                    "gl_staff_balance": float(employee_gl["staff_balance"]) if employee_gl else 0,
                    "gl_unearned_balance": float(employee_gl["unearned_balance"]) if employee_gl else 0,
                    "gl_interest_earned": float(employee_gl["interest_earned"]) if employee_gl else 0,
                    "gl_total_outstanding": float(employee_gl["total_outstanding"]) if employee_gl else 0,
                    "source_opening_count": employee_gl["opening_count"] if employee_gl else 0,
                    "source_repayment_count": employee_gl["repayment_count"] if employee_gl else 0,
                    "source_vouchers": "\n".join(employee_gl["source_vouchers"]) if employee_gl else "",
                    "validation_status": status,
                    "validation_message": message,
                },
            )

        self.update_counts()
        self.save(ignore_permissions=True)
        return counts_dict(self)

    def before_submit(self):
        if not self.rows:
            frappe.throw(_("Load and validate the outstanding report before submitting"))
        if not self.confirm_no_bank_posting:
            frappe.throw(_("Confirm that the batch must not post to Bank or Cash"))
        self.update_counts()
        if not self.ready_rows:
            frappe.throw(_("The batch has no rows ready for conversion"))
        self.status = "Queued"

    def on_submit(self):
        frappe.enqueue(
            "employee_lending.employee_lending.doctype.employee_loan_legacy_import_batch.employee_loan_legacy_import_batch.process_batch",
            queue="long",
            timeout=7200,
            enqueue_after_commit=True,
            batch_name=self.name,
        )

    def before_cancel(self):
        processed = frappe.db.exists(
            "Employee Loan Legacy Import Row",
            {"parent": self.name, "validation_status": "Processed"},
        )
        if processed:
            frappe.throw(
                _("This batch has processed loans. Cancel the linked Legacy Loan Openings before cancelling the batch")
            )

    def on_cancel(self):
        self.db_set("status", "Cancelled", update_modified=False)

    def update_counts(self):
        statuses = [row.validation_status for row in self.rows]
        self.total_rows = len(statuses)
        self.ready_rows = statuses.count("Ready")
        self.excluded_rows = statuses.count("Excluded")
        self.invalid_rows = statuses.count("Invalid")
        self.processed_rows = statuses.count("Processed")
        self.failed_rows = statuses.count("Failed")
        ready = [row for row in self.rows if row.validation_status == "Ready"]
        self.ready_principal_outstanding = sum(flt(row.principal_outstanding) for row in ready)
        self.ready_interest_outstanding = sum(flt(row.interest_outstanding) for row in ready)
        self.ready_total_outstanding = sum(flt(row.reported_total_outstanding) for row in ready)
        self.source_opening_vouchers = sum(int(row.source_opening_count or 0) for row in self.rows)
        self.source_repayment_allocations = sum(int(row.source_repayment_count or 0) for row in self.rows)


def read_outstanding_file(file_url):
    try:
        from openpyxl import load_workbook
    except ImportError:
        frappe.throw(_("The server is missing the openpyxl package required to read XLSX files"))

    file_doc = frappe.get_doc("File", {"file_url": file_url})
    file_path = file_doc.get_full_path()
    if not os.path.isfile(file_path):
        frappe.throw(_("Attached source file was not found on the server"))

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    header_row = None
    headers = None
    for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        try:
            candidate = map_headers(values)
        except ValueError:
            continue
        header_row, headers = row_number, candidate
        break
    if not headers:
        frappe.throw(
            _("This is not the Employee Loan Outstanding report. Required outstanding columns were not found")
        )

    records = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        try:
            record = classify_outstanding_row(values, headers)
        except ValueError as exc:
            employee_index = headers.get("employee")
            employee = str(values[employee_index] or "").strip() if employee_index is not None else ""
            if not employee:
                continue
            record = empty_invalid_record(employee, str(exc))
        records.append((row_number, record))
    workbook.close()
    return records


def read_gl_file(file_url, settings, cutoff_date):
    workbook, sheet = open_xlsx(file_url)
    header_row = None
    headers = None
    for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        try:
            candidate = map_gl_headers(values)
        except ValueError:
            continue
        header_row, headers = row_number, candidate
        break
    if not headers:
        workbook.close()
        frappe.throw(_("This is not the required General Ledger export. GL columns were not found"))

    rows = []
    cutoff = parse_source_date(cutoff_date)
    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        voucher_index = headers["voucher_no"]
        if not values[voucher_index]:
            continue
        posting_date = parse_source_date(values[headers["posting_date"]])
        if posting_date and posting_date > cutoff:
            workbook.close()
            frappe.throw(_("The General Ledger export contains entries after the selected cut-off date"))
        rows.append(
            {
                fieldname: values[index]
                for fieldname, index in headers.items()
            }
        )
        rows[-1]["posting_date"] = posting_date.isoformat() if posting_date else ""
    workbook.close()
    return aggregate_gl_rows(
        rows,
        settings.staff_loan_receivable_account,
        settings.unearned_interest_account,
        settings.interest_income_account,
        settings.legacy_temporary_account,
    )


def open_xlsx(file_url):
    try:
        from openpyxl import load_workbook
    except ImportError:
        frappe.throw(_("The server is missing the openpyxl package required to read XLSX files"))
    file_doc = frappe.get_doc("File", {"file_url": file_url})
    file_path = file_doc.get_full_path()
    if not os.path.isfile(file_path):
        frappe.throw(_("Attached source file was not found on the server"))
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    return workbook, workbook[workbook.sheetnames[0]]


def parse_source_date(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for date_format in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    frappe.throw(_("Invalid posting date in source file: {0}").format(text))


def validate_gl_reconciliation(record, employee_gl):
    if not employee_gl:
        return "Invalid", "Employee has no matching Journal Entry activity in the GL export"
    if not employee_gl["opening_count"]:
        return "Invalid", "No opening or disbursement Journal Entry was identified"
    checks = (
        (employee_gl["staff_balance"], record["principal_outstanding"], "Staff Loan balance"),
        (
            employee_gl["unearned_balance"] - employee_gl["interest_earned"],
            record["interest_outstanding"],
            "interest balance",
        ),
        (employee_gl["total_outstanding"], record["reported_total_outstanding"], "total outstanding"),
    )
    for gl_value, report_value, label in checks:
        if abs(float(gl_value) - float(report_value)) > 0.01:
            return "Invalid", f"GL {label} does not match the outstanding report"
    if employee_gl["opening_count"] > 1:
        return "Ready", f"Consolidating {employee_gl['opening_count']} historical opening/disbursement vouchers"
    return "Ready", ""


def empty_invalid_record(employee, message):
    return {
        "employee": employee,
        "original_principal": 0,
        "original_interest": 0,
        "principal_repaid": 0,
        "interest_repaid": 0,
        "principal_outstanding": 0,
        "interest_outstanding": 0,
        "reported_total_outstanding": 0,
        "calculated_total_outstanding": 0,
        "source_interest_rate": 0,
        "source_remarks": "",
        "validation_status": "Invalid",
        "validation_message": message,
    }


def match_product(source_rate, products):
    matches = [product for product in products if abs(flt(product.flat_interest_rate) - flt(source_rate)) <= 0.25]
    if len(matches) != 1:
        return None
    return matches[0]


def counts_dict(batch):
    return {
        "total_rows": batch.total_rows,
        "ready_rows": batch.ready_rows,
        "excluded_rows": batch.excluded_rows,
        "invalid_rows": batch.invalid_rows,
        "processed_rows": batch.processed_rows,
        "failed_rows": batch.failed_rows,
        "source_opening_vouchers": batch.source_opening_vouchers,
        "source_repayment_allocations": batch.source_repayment_allocations,
    }


@frappe.whitelist()
def retry_failed_rows(batch_name):
    batch = frappe.get_doc("Employee Loan Legacy Import Batch", batch_name)
    batch.check_permission("submit")
    if batch.docstatus != 1 or batch.status in ("Queued", "Running"):
        frappe.throw(_("This batch cannot be retried in its current state"))
    frappe.db.sql(
        """
        update `tabEmployee Loan Legacy Import Row`
           set validation_status = 'Ready', validation_message = ''
         where parent = %s and validation_status = 'Failed'
        """,
        batch.name,
    )
    frappe.db.set_value("Employee Loan Legacy Import Batch", batch.name, "status", "Queued")
    frappe.enqueue(
        "employee_lending.employee_lending.doctype.employee_loan_legacy_import_batch.employee_loan_legacy_import_batch.process_batch",
        queue="long",
        timeout=7200,
        enqueue_after_commit=True,
        batch_name=batch.name,
    )
    return {"status": "Queued"}


def process_batch(batch_name):
    batch = frappe.get_doc("Employee Loan Legacy Import Batch", batch_name)
    if batch.docstatus != 1 or batch.status not in ("Queued", "Running"):
        return

    frappe.db.set_value(
        "Employee Loan Legacy Import Batch",
        batch.name,
        {"status": "Running", "started_on": batch.started_on or now(), "completed_on": None},
    )
    frappe.db.commit()

    settings = get_settings()
    gl_by_employee = read_gl_file(batch.gl_source_file, settings, batch.cutoff_date)

    rows = frappe.get_all(
        "Employee Loan Legacy Import Row",
        filters={"parent": batch.name, "validation_status": "Ready"},
        fields=["*"],
        order_by="idx",
    )
    for row in rows:
        frappe.db.set_value("Employee Loan Legacy Import Row", row.name, "validation_status", "Processing")
        frappe.db.commit()
        try:
            opening = create_legacy_opening(batch, row)
            employee_gl = gl_by_employee.get(row.employee)
            if not employee_gl:
                frappe.throw(_("No matching source GL history was found for employee {0}").format(row.employee))
            create_legacy_transactions(batch, row, opening, employee_gl["transactions"])
            frappe.db.set_value(
                "Employee Loan Legacy Import Row",
                row.name,
                {
                    "validation_status": "Processed",
                    "validation_message": "",
                    "legacy_opening": opening.name,
                },
            )
            frappe.db.commit()
        except Exception as exc:
            frappe.db.rollback()
            error = frappe.get_traceback(with_context=False)
            message = str(exc) or (error.splitlines()[-1] if error else "Unknown conversion error")
            frappe.db.set_value(
                "Employee Loan Legacy Import Row",
                row.name,
                {"validation_status": "Failed", "validation_message": message[:500]},
            )
            frappe.log_error(title=f"Legacy loan batch {batch.name}, row {row.row_number}", message=error)
            frappe.db.commit()
        finally:
            if getattr(frappe.local, "message_log", None):
                frappe.local.message_log = []

    update_batch_completion(batch.name)


def create_legacy_opening(batch, row):
    opening = frappe.get_doc(
        {
            "doctype": "Employee Loan Legacy Opening",
            "cutoff_date": batch.cutoff_date,
            "source_reference": (f"{batch.source_reference}; {batch.name}; row {row.row_number}")[:140],
            "employee": row.employee,
            "loan_product": row.loan_product,
            "original_principal": row.original_principal,
            "original_interest": row.original_interest,
            "reported_total_outstanding": row.reported_total_outstanding,
            "principal_repaid": row.principal_repaid,
            "interest_repaid": row.interest_repaid,
            "next_repayment_date": batch.next_repayment_date,
            "confirm_no_bank_posting": 1,
            "remarks": f"Created from bulk legacy import {batch.name}, source row {row.row_number}",
        }
    )
    opening.insert(ignore_permissions=True)
    opening.flags.ignore_permissions = True
    opening.submit()
    return opening


def create_legacy_transactions(batch, row, opening, transactions):
    loan = frappe.get_doc("Employee Loan Application", opening.loan_application)
    for transaction in transactions:
        source_key = "{0}|{1}|{2}".format(batch.name, row.employee, transaction["voucher"])
        history = frappe.get_doc(
            {
                "doctype": "Employee Loan Legacy Transaction",
                "source_key": source_key[:140],
                "legacy_import_batch": batch.name,
                "legacy_opening": opening.name,
                "loan_application": loan.name,
                "employee": row.employee,
                "employee_name": loan.employee_name,
                "posting_date": transaction["posting_date"],
                "transaction_type": transaction["transaction_type"],
                "journal_entry": transaction["voucher"],
                "principal_amount": float(transaction["principal_amount"]),
                "interest_amount": float(transaction["interest_amount"]),
                "debit": float(transaction["debit"]),
                "credit": float(transaction["credit"]),
                "against_account": transaction["against_account"],
                "remarks": _("Existing production Journal Entry; linked by bulk legacy import {0}").format(
                    batch.name
                ),
            }
        )
        history.flags.from_legacy_import = True
        history.insert(ignore_permissions=True)


def update_batch_completion(batch_name):
    counts = dict(
        frappe.db.sql(
            """
            select validation_status, count(*)
              from `tabEmployee Loan Legacy Import Row`
             where parent = %s
             group by validation_status
            """,
            batch_name,
        )
    )
    failed = counts.get("Failed", 0)
    invalid = counts.get("Invalid", 0)
    values = {
        "ready_rows": counts.get("Ready", 0),
        "excluded_rows": counts.get("Excluded", 0),
        "invalid_rows": invalid,
        "processed_rows": counts.get("Processed", 0),
        "failed_rows": failed,
        "status": "Completed with Errors" if failed or invalid else "Completed",
        "completed_on": now(),
    }
    frappe.db.set_value("Employee Loan Legacy Import Batch", batch_name, values)
    frappe.db.commit()
